import json
import sys
from pathlib import Path

files = [
    r"c:\Users\DELL\Downloads\Gorsel_Programlama_II__Ders_1.pdf",
    r"c:\Users\DELL\Downloads\Cerezler__Oturumlar_ve_Passport_Kimlik_.pdf",
    r"c:\Users\DELL\Downloads\PostgreSQL_e_Giris.pdf",
    r"c:\Users\DELL\Downloads\API_ler.pptx",
    r"c:\Users\DELL\Downloads\Ders_7_VeriTabani.pdf",
    r"c:\Users\DELL\Downloads\Express_Middleware_ve_Body-Parser_DERS_5 (1).pptx",
    r"c:\Users\DELL\Downloads\Gorsel_Programa_Ders_3 (1).pptx",
    r"c:\Users\DELL\Downloads\Gorsel_Programlama_II_Ders_2 (1).pptx",
]

results = {}

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

from pptx import Presentation

def extract_pdf(path):
    text_parts = []
    if pdfplumber:
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages):
                t = page.extract_text() or ""
                text_parts.append(f"--- Sayfa {i+1} ---\n{t}")
    else:
        from pypdf import PdfReader
        reader = PdfReader(path)
        for i, page in enumerate(reader.pages):
            t = page.extract_text() or ""
            text_parts.append(f"--- Sayfa {i+1} ---\n{t}")
    return "\n\n".join(text_parts)

def extract_pptx(path):
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        if slide_text:
            parts.append(f"--- Slayt {i+1} ---\n" + "\n".join(slide_text))
    return "\n\n".join(parts)

for f in files:
    p = Path(f)
    if not p.exists():
        results[p.name] = {"error": "file not found"}
        continue
    try:
        if p.suffix.lower() == ".pdf":
            results[p.name] = {"type": "pdf", "text": extract_pdf(str(p))}
        elif p.suffix.lower() in (".pptx", ".ppt"):
            results[p.name] = {"type": "pptx", "text": extract_pptx(str(p))}
    except Exception as e:
        results[p.name] = {"error": str(e)}

out = Path(r"c:\Users\DELL\Desktop\ders\extracted_content.json")
out.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Extracted {len(results)} files -> {out}")
for name, data in results.items():
    if "text" in data:
        print(f"{name}: {len(data['text'])} chars")
    else:
        print(f"{name}: ERROR - {data.get('error')}")
