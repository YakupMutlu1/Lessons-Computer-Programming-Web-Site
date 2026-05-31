# -*- coding: utf-8 -*-
import json
from pathlib import Path

files = [
    r"c:\Users\DELL\Downloads\Hafta1_Internet_Guvenligine_Giris.pdf",
    r"c:\Users\DELL\Downloads\Hafta2_Ag_Guvenligi_ve_Ag_Katmanlari.pdf",
    r"c:\Users\DELL\Downloads\Hafta3_Sifreleme_ve_Kriptografi_Temelleri.pdf",
    r"c:\Users\DELL\Downloads\Hafta4_Kimlik_Dogrulama_ve_Erisim_Yonetimi.pdf",
    r"c:\Users\DELL\Downloads\Hafta5_Zararli_Yazilimlar.pdf",
    r"c:\Users\DELL\Downloads\Hafta6_Sosyal_Muhendislik_Saldirilari.pdf",
    r"c:\Users\DELL\Downloads\Hafta9_Mobil_Guvenlik_IoT_Guvenligi.pdf",
    r"c:\Users\DELL\Downloads\Hafta10_Ag_Izleme_ve_Guvenlik_Testleri.pdf",
    r"c:\Users\DELL\Downloads\Hafta11_Guvenlik_Protokolleri_ve_Standartlari.pdf",
    r"c:\Users\DELL\Downloads\Hafta12_Siber_Tehditler_ve_Acil_Durum_Yonetimi.pdf",
    r"c:\Users\DELL\Downloads\Hafta14_Internet_Etigi.pptx",
]

import pdfplumber
from pptx import Presentation

out_dir = Path(r"c:\Users\DELL\Desktop\ders\docs\guvenlik")
out_dir.mkdir(parents=True, exist_ok=True)
results = {}

for f in files:
    p = Path(f)
    if not p.exists():
        results[p.name] = "NOT FOUND"
        continue
    try:
        if p.suffix.lower() == ".pdf":
            parts = []
            with pdfplumber.open(p) as pdf:
                for i, page in enumerate(pdf.pages):
                    parts.append(f"--- Sayfa {i+1} ---\n{page.extract_text() or ''}")
            text = "\n\n".join(parts)
        else:
            parts = []
            prs = Presentation(p)
            for i, slide in enumerate(prs.slides):
                t = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        t.append(shape.text.strip())
                if t:
                    parts.append(f"--- Slayt {i+1} ---\n" + "\n".join(t))
            text = "\n\n".join(parts)
        safe = p.stem
        (out_dir / f"{safe}.txt").write_text(text, encoding="utf-8")
        results[p.name] = len(text)
        print(f"{p.name}: {len(text)} chars")
    except Exception as e:
        results[p.name] = str(e)
        print(f"{p.name}: ERROR {e}")

Path(r"c:\Users\DELL\Desktop\ders\docs\guvenlik\_index.json").write_text(
    json.dumps(results, indent=2), encoding="utf-8"
)
